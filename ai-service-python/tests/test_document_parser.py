"""document_parser 单元测试（Wave B #5）"""

import io
import os
from pathlib import Path

import pytest
from docx import Document
from fastapi import UploadFile

from app.services.binary_document_extractor import DocumentExtractError, extract_docx_text
from app.services.document_parser import ALLOWED_EXTENSIONS, parse_upload_file


def _upload(name: str, data: bytes) -> UploadFile:
    return UploadFile(filename=name, file=io.BytesIO(data))


def test_allowed_extensions_include_pdf_docx():
    assert ".pdf" in ALLOWED_EXTENSIONS
    assert ".docx" in ALLOWED_EXTENSIONS


def test_parse_utf8_markdown():
    raw = "# Title\n\nhello".encode("utf-8")
    filename, content = parse_upload_file(_upload("note.md", raw))
    assert filename == "note.md"
    assert "hello" in content


def test_parse_docx_paragraphs():
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Wave B RAG paragraph")
    doc.save(buf)
    text = extract_docx_text(buf.getvalue())
    assert "Wave B RAG paragraph" in text


def test_parse_empty_docx_raises():
    buf = io.BytesIO()
    Document().save(buf)
    with pytest.raises(DocumentExtractError):
        extract_docx_text(buf.getvalue())


def test_unsupported_extension():
    with pytest.raises(Exception) as exc:
        parse_upload_file(_upload("data.xlsx", b"123"))
    assert exc.value.status_code == 400


@pytest.mark.skipif(
    not os.getenv("DOCX_FIXTURE_PATH"),
    reason="Set DOCX_FIXTURE_PATH to a local textbox-layout DOCX for manual regression",
)
def test_extract_docx_from_local_fixture():
    fixture_path = Path(os.environ["DOCX_FIXTURE_PATH"])
    if not fixture_path.is_file():
        pytest.skip(f"DOCX fixture not found: {fixture_path}")

    text = extract_docx_text(fixture_path.read_bytes())
    assert len(text) > 100
